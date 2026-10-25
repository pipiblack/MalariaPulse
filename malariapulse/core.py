"""Deterministic malaria surveillance, response planning, and local persistence."""

from __future__ import annotations

import csv
import io
import json
import math
import os
import re
from copy import deepcopy
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
SEED_PATH = ROOT / "data" / "seed.json"
OPERATIONAL_PATH = ROOT / "data" / "operational.json"
RUNTIME_PATH = ROOT / "data" / "runtime.json"

RESOURCE_LABELS = {"rdt": "RDT kits", "act": "ACT courses", "nets": "treated nets"}
MAX_IMPORT_BYTES = 15 * 1024 * 1024

IMPORT_SCHEMAS = {
    "clinical": {
        "id_field": "zone_id",
        "required": [
            "zone_id",
            "reporting_week",
            "week_1_cases",
            "week_2_cases",
            "week_3_cases",
            "week_4_cases",
            "week_5_cases",
            "week_6_cases",
            "tests",
            "positive_tests",
        ],
        "optional": [
            "suspected_cases",
            "severe_cases",
            "admissions",
            "deaths",
            "under_5_confirmed_cases",
            "pregnancy_confirmed_cases",
            "chw_fever_referrals",
            "reporting_completeness_pct",
        ],
    },
    "environment": {
        "id_field": "zone_id",
        "required": ["zone_id", "rainfall_7d_mm", "temperature_c", "humidity_pct"],
        "optional": ["rainfall_30d_mm", "standing_water_index", "flooded_sites"],
    },
    "intervention": {
        "id_field": "zone_id",
        "required": [
            "zone_id",
            "coverage_date",
            "llin_use_pct",
            "iptp3_coverage_pct",
            "vaccine_dose4_coverage_pct",
            "larval_sites_managed_pct",
        ],
    },
    "supply": {
        "id_field": "clinic_id",
        "required": [
            "clinic_id",
            "stock_rdt",
            "stock_act",
            "stock_nets",
            "weekly_use_rdt",
            "weekly_use_act",
            "weekly_use_nets",
        ],
    },
}

SENSITIVE_COLUMNS = {
    "patient_id",
    "patient_name",
    "full_name",
    "first_name",
    "last_name",
    "phone",
    "phone_number",
    "national_id",
    "id_number",
    "date_of_birth",
    "dob",
    "address",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z")


def _is_iso_date(value: str) -> bool:
    try:
        return datetime.strptime(value, "%Y-%m-%d").strftime("%Y-%m-%d") == value
    except ValueError:
        return False


def _read_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def _atomic_write(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    os.replace(temporary, path)


def load_seed_data() -> dict[str, Any]:
    return _with_operational_defaults(_read_json(SEED_PATH))


def load_dataset() -> dict[str, Any]:
    return _with_operational_defaults(_read_json(OPERATIONAL_PATH if OPERATIONAL_PATH.exists() else SEED_PATH))


def _with_operational_defaults(dataset: dict[str, Any]) -> dict[str, Any]:
    """Keep older saved datasets compatible as the operational indicator set grows."""

    metadata_completeness = float(dataset.get("metadata", {}).get("dataCompleteness", 0.9)) * 100
    for zone in dataset.get("zones", []):
        tests = int(zone.get("tests", 0))
        clinical = zone.setdefault("clinical", {})
        clinical.setdefault("suspectedCases", tests)
        clinical.setdefault("severeCases", 0)
        clinical.setdefault("admissions", 0)
        clinical.setdefault("deaths", 0)
        clinical.setdefault("under5ConfirmedCases", 0)
        clinical.setdefault("pregnancyConfirmedCases", 0)
        clinical.setdefault("chwFeverReferrals", 0)
        clinical.setdefault("reportingCompletenessPct", round(metadata_completeness, 1))

        environment = zone.setdefault("environment", {})
        rainfall_7d = float(environment.get("rainfall7dMm", 0))
        environment.setdefault("rainfall30dMm", round(rainfall_7d * 3.2, 1))
        environment.setdefault("standingWaterIndex", 0)
        environment.setdefault("floodedSites", 0)

        zone.setdefault(
            "interventions",
            {
                "coverageDate": dataset.get("metadata", {}).get("updatedAt", "")[:10],
                "llinUsePct": 0,
                "iptp3CoveragePct": 0,
                "vaccineDose4CoveragePct": 0,
                "larvalSitesManagedPct": 0,
            },
        )
    return dataset


def save_dataset(dataset: dict[str, Any]) -> None:
    _atomic_write(OPERATIONAL_PATH, dataset)


def _clamp(value: float, minimum: float = 0, maximum: float = 1) -> float:
    return min(maximum, max(minimum, value))


def _average(values: list[float]) -> float:
    return sum(values) / max(len(values), 1)


def _round(value: float, precision: int = 0) -> int | float:
    factor = 10**precision
    rounded = math.floor(value * factor + 0.5) / factor
    return int(rounded) if precision == 0 else rounded


def classify_risk(score: int) -> str:
    if score >= 75:
        return "critical"
    if score >= 60:
        return "high"
    if score >= 40:
        return "watch"
    return "low"


def assess_zone(zone: dict[str, Any], metadata: dict[str, Any] | None = None) -> dict[str, Any]:
    history = zone.get("weeklyCases") if zone else None
    if not isinstance(history, list) or len(history) < 4:
        raise ValueError("A zone with at least four weekly case counts is required.")

    baseline = _average(history[:-2])
    recent = _average(history[-2:])
    latest = history[-1]
    growth_ratio = recent / max(baseline, 1)
    positivity = zone["positiveTests"] / max(zone["tests"], 1)
    environment = zone["environment"]
    clinical = zone.get("clinical", {})
    interventions = zone.get("interventions", {})
    suspected_cases = max(int(clinical.get("suspectedCases", zone["tests"])), 0)
    reporting_completeness_pct = float(
        clinical.get("reportingCompletenessPct", (metadata or {}).get("dataCompleteness", 0.9) * 100)
    )
    testing_coverage_pct = _round(zone["tests"] / max(suspected_cases, 1) * 100, 1) if suspected_cases else None
    components = [
        {
            "key": "caseTrend",
            "label": f"{_round((growth_ratio - 1) * 100)}% case acceleration",
            "value": _clamp((growth_ratio - 0.9) / 1.25),
            "weight": 0.4,
        },
        {
            "key": "positivity",
            "label": f"{_round(positivity * 100)}% test positivity",
            "value": _clamp(positivity / 0.34),
            "weight": 0.23,
        },
        {
            "key": "rainfall",
            "label": f"{environment['rainfall7dMm']} mm rainfall / 7 days",
            "value": _clamp(environment["rainfall7dMm"] / 130),
            "weight": 0.16,
        },
        {
            "key": "humidity",
            "label": f"{environment['humidityPct']}% humidity",
            "value": _clamp((environment["humidityPct"] - 52) / 30),
            "weight": 0.11,
        },
        {
            "key": "temperature",
            "label": f"{environment['temperatureC']}°C transmission suitability",
            "value": _clamp(1 - abs(environment["temperatureC"] - 27) / 9),
            "weight": 0.1,
        },
    ]
    score = int(_round(sum(item["value"] * item["weight"] for item in components) * 100))
    risk_level = classify_risk(score)
    projected_multiplier = _clamp(0.6 + growth_ratio * 0.55, 0.75, 1.85)
    data_completeness = _clamp(reporting_completeness_pct / 100)
    confidence = _round(_clamp(0.7 + data_completeness * 0.24 - (0.15 if zone["tests"] < 30 else 0)), 2)
    recommendation = {
        "critical": "Activate rapid response review within 4 hours.",
        "high": "Validate the alert and prepare targeted response within 24 hours.",
        "watch": "Increase reporting cadence and verify case data.",
        "low": "Continue routine surveillance.",
    }[risk_level]

    drivers = sorted(
        [
            {"label": item["label"], "contribution": _round(item["value"] * item["weight"] * 100, 1)}
            for item in components
        ],
        key=lambda item: item["contribution"],
        reverse=True,
    )[:3]
    priority_flags: list[str] = []
    if int(clinical.get("deaths", 0)) > 0:
        priority_flags.append("Verify reported malaria death immediately")
    if int(clinical.get("severeCases", 0)) >= 5:
        priority_flags.append("Severe-case burden requires referral-readiness review")
    if testing_coverage_pct is not None and testing_coverage_pct < 90:
        priority_flags.append("Fewer than 90% of suspected cases were tested")
    if reporting_completeness_pct < 90:
        priority_flags.append("Reporting completeness is below 90%")
    if float(environment.get("standingWaterIndex", 0)) >= 75:
        priority_flags.append("High standing-water observation")
    if 0 < float(interventions.get("llinUsePct", 0)) < 80:
        priority_flags.append("Reported LLIN use is below 80%")
    return {
        "zoneId": zone["id"],
        "zoneName": zone["name"],
        "subcounty": zone["subcounty"],
        "score": score,
        "riskLevel": risk_level,
        "confidence": confidence,
        "currentCases": latest,
        "expectedCasesNextWeek": int(_round(latest * projected_multiplier)),
        "growthPct": _round((growth_ratio - 1) * 100),
        "positivityPct": _round(positivity * 100, 1),
        "testingCoveragePct": testing_coverage_pct,
        "incidencePerThousand": _round((latest / zone["population"]) * 1000, 2),
        "coordinates": zone["coordinates"],
        "drivers": drivers,
        "priorityFlags": priority_flags,
        "clinicalContext": {
            "suspectedCases": suspected_cases,
            "severeCases": int(clinical.get("severeCases", 0)),
            "admissions": int(clinical.get("admissions", 0)),
            "deaths": int(clinical.get("deaths", 0)),
            "under5ConfirmedCases": int(clinical.get("under5ConfirmedCases", 0)),
            "pregnancyConfirmedCases": int(clinical.get("pregnancyConfirmedCases", 0)),
            "chwFeverReferrals": int(clinical.get("chwFeverReferrals", 0)),
            "reportingCompletenessPct": _round(reporting_completeness_pct, 1),
        },
        "environmentContext": {
            "rainfall30dMm": environment.get("rainfall30dMm"),
            "standingWaterIndex": environment.get("standingWaterIndex"),
            "floodedSites": environment.get("floodedSites"),
        },
        "interventionContext": interventions,
        "recommendation": recommendation,
    }


def _distance_km(first: dict[str, float], second: dict[str, float]) -> float:
    radius = 6371
    lat1, lat2 = math.radians(first["lat"]), math.radians(second["lat"])
    d_lat = math.radians(second["lat"] - first["lat"])
    d_lng = math.radians(second["lng"] - first["lng"])
    haversine = math.sin(d_lat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(d_lng / 2) ** 2
    return radius * 2 * math.atan2(math.sqrt(haversine), math.sqrt(1 - haversine))


def plan_resources(dataset: dict[str, Any], zone_id: str, assessment: dict[str, Any]) -> dict[str, Any]:
    working = deepcopy(dataset)
    zone = next((item for item in working["zones"] if item["id"] == zone_id), None)
    if not zone:
        raise ValueError(f"Unknown zone: {zone_id}")
    target_clinics = [clinic for clinic in working["clinics"] if clinic["id"] in zone["clinicIds"]]
    donor_clinics = [clinic for clinic in working["clinics"] if clinic["id"] not in zone["clinicIds"]]
    multiplier = {"critical": 1.7, "high": 1.45}.get(assessment["riskLevel"], 1.2)
    recommendations: list[dict[str, Any]] = []
    restock_requests: list[dict[str, Any]] = []

    for target in target_clinics:
        for resource, label in RESOURCE_LABELS.items():
            target_level = math.ceil(target["weeklyUse"][resource] * 3 * multiplier)
            deficit = max(0, target_level - target["stock"][resource])
            if deficit == 0:
                continue
            donors = sorted(
                [
                    {
                        "clinic": clinic,
                        "surplus": max(0, clinic["stock"][resource] - math.ceil(clinic["weeklyUse"][resource] * 3)),
                        "distance": _distance_km(clinic["coordinates"], target["coordinates"]),
                    }
                    for clinic in donor_clinics
                ],
                key=lambda candidate: candidate["distance"],
            )
            for donor in [candidate for candidate in donors if candidate["surplus"] > 0]:
                if deficit <= 0:
                    break
                quantity = min(deficit, donor["surplus"])
                receiver_stock_before = target["stock"][resource]
                donor_stock_before = donor["clinic"]["stock"][resource]
                receiver_use = target["weeklyUse"][resource]
                donor_use = donor["clinic"]["weeklyUse"][resource]
                target["stock"][resource] += quantity
                donor["clinic"]["stock"][resource] -= quantity
                recommendations.append(
                    {
                        "priority": "urgent" if assessment["riskLevel"] == "critical" else "high",
                        "resource": resource,
                        "resourceLabel": label,
                        "quantity": quantity,
                        "fromClinicId": donor["clinic"]["id"],
                        "fromClinic": donor["clinic"]["name"],
                        "toClinicId": target["id"],
                        "toClinic": target["name"],
                        "distanceKm": _round(donor["distance"], 1),
                        "etaMinutes": max(25, int(_round(donor["distance"] * 2.4 + 18))),
                        "receiverWeeksBefore": _round(receiver_stock_before / max(receiver_use, 1), 1),
                        "receiverWeeksAfter": _round(target["stock"][resource] / max(receiver_use, 1), 1),
                        "donorWeeksAfter": _round((donor_stock_before - quantity) / max(donor_use, 1), 1),
                        "targetWeeks": _round(3 * multiplier, 1),
                        "rationale": f"Restores approximately three weeks of safety stock at projected {assessment['riskLevel']} demand.",
                    }
                )
                deficit -= quantity
            if deficit > 0:
                restock_requests.append(
                    {
                        "resource": resource,
                        "resourceLabel": label,
                        "quantity": deficit,
                        "clinicId": target["id"],
                        "clinic": target["name"],
                        "priority": "urgent" if assessment["riskLevel"] == "critical" else "high",
                        "stockOnHand": target["stock"][resource],
                        "weeklyUse": target["weeklyUse"][resource],
                        "coverageWeeks": _round(
                            target["stock"][resource] / max(target["weeklyUse"][resource], 1), 1
                        ),
                        "targetStock": target_level,
                    }
                )

    readiness_values = [
        _clamp(clinic["stock"][resource] / max(clinic["weeklyUse"][resource] * 3 * multiplier, 1))
        for clinic in target_clinics
        for resource in RESOURCE_LABELS
    ]
    return {
        "zoneId": zone_id,
        "generatedForRisk": assessment["riskLevel"],
        "readinessPct": int(_round(_average(readiness_values) * 100)),
        "recommendations": recommendations,
        "restockRequests": restock_requests,
        "requiresApproval": bool(recommendations or restock_requests),
        "method": (
            "The draft protects three weeks of stock at donor facilities, adjusts the receiving target for the current "
            f"{assessment['riskLevel']} risk level, and prioritizes shorter estimated routes."
        ),
        "note": "Draft only. A county pharmacist or authorized logistics officer must verify counts and approve movement.",
    }


def draft_advisories(dataset: dict[str, Any], zone_id: str, assessment: dict[str, Any]) -> dict[str, Any]:
    zone = next((item for item in dataset["zones"] if item["id"] == zone_id), None)
    if not zone:
        raise ValueError(f"Unknown zone: {zone_id}")
    clinics = [clinic for clinic in dataset["clinics"] if clinic["id"] in zone["clinicIds"]]
    clinic_name = clinics[0]["name"] if clinics else "your nearest health facility"
    english = (
        f"Malaria risk is rising in {zone['name']}. Sleep under a treated net every night. If you have fever, chills or "
        f"headache, seek a malaria test promptly at {clinic_name}. Children under 5 and pregnant people need prompt "
        "attention. Test before treatment. This is an alert, not a diagnosis."
    )
    swahili = (
        f"Hatari ya malaria inaongezeka {zone['name']}. Lala ndani ya chandarua chenye dawa kila usiku. Ukipata homa, "
        f"baridi au maumivu ya kichwa, nenda kupimwa mapema katika {clinic_name}. Watoto chini ya miaka 5 na wajawazito "
        "wahudumiwe haraka. Pima kabla ya dawa. Huu ni ujumbe wa tahadhari, si utambuzi."
    )
    chw = (
        f"CHW ACTION — {zone['name']}: verify fever reports, encourage prompt testing, check treated-net use, and refer "
        "severe illness immediately. Prioritize children under 5 and pregnant people. Do not recommend antimalarials "
        "without an approved test-and-treat protocol. Report stock gaps and unusual clusters to the sub-county surveillance lead."
    )

    reason = (
        f"Prepared because {zone['name']} has a {assessment['riskLevel']} operational signal "
        f"({assessment['score']}/100), {assessment['growthPct']}% case acceleration, and "
        f"{assessment['positivityPct']}% test positivity."
    )

    def message(
        audience: str,
        language: str,
        channel: str,
        text: str,
        purpose: str,
        delivery_notes: str,
    ) -> dict[str, Any]:
        return {
            "audience": audience,
            "language": language,
            "channel": channel,
            "text": text,
            "purpose": purpose,
            "reason": reason,
            "deliveryNotes": delivery_notes,
            "reviewChecklist": [
                "Confirm the named facility is open and able to test suspected malaria cases.",
                "Confirm the message is limited to the intended micro-zone and approved sender list.",
                "Have an authorized county reviewer approve wording, timing, and release channel.",
            ],
            "characters": len(text),
            "segments": math.ceil(len(text) / 153),
        }

    return {
        "zoneId": zone_id,
        "messages": [
            message(
                "residents",
                "English",
                "SMS/WhatsApp",
                english,
                "Encourage nightly LLIN use and prompt testing without declaring an outbreak or diagnosis.",
                "Use the county-approved sender ID after facility capacity and referral details are verified.",
            ),
            message(
                "residents",
                "Kiswahili",
                "SMS/WhatsApp",
                swahili,
                "Provide the same prevention and prompt-testing guidance in Kiswahili.",
                "Have a fluent local reviewer confirm wording before release and pair it with verified facility details.",
            ),
            message(
                "community health workers",
                "English",
                "WhatsApp",
                chw,
                "Give CHWs a verification, referral, prevention, and stock-escalation checklist.",
                "Release through the authorized CHW coordination group with the current sub-county escalation contact.",
            ),
        ],
        "requiresApproval": True,
        "summary": reason,
        "clinicalSafety": "Templates follow prevention and prompt-testing principles; local health authorities must review before release.",
    }


def build_network_snapshot(dataset: dict[str, Any] | None = None) -> dict[str, Any]:
    dataset = dataset or load_dataset()
    assessments = sorted(
        [assess_zone(zone, dataset["metadata"]) for zone in dataset["zones"]], key=lambda item: item["score"], reverse=True
    )
    stock = {
        resource: sum(clinic["stock"][resource] for clinic in dataset["clinics"]) for resource in RESOURCE_LABELS
    }
    readiness = [
        _average(
            [
                _clamp(clinic["stock"][resource] / max(clinic["weeklyUse"][resource] * 3, 1))
                for resource in RESOURCE_LABELS
            ]
        )
        for clinic in dataset["clinics"]
    ]
    latest_cases = sum(item["currentCases"] for item in assessments)
    previous_cases = sum(zone["weeklyCases"][-2] for zone in dataset["zones"])
    tests = sum(zone["tests"] for zone in dataset["zones"])
    positives = sum(zone["positiveTests"] for zone in dataset["zones"])
    clinical_context = [item["clinicalContext"] for item in assessments]
    facilities_below_three_weeks = sum(
        any(
            clinic["weeklyUse"][resource] > 0
            and clinic["stock"][resource] / clinic["weeklyUse"][resource] < 3
            for resource in ("rdt", "act", "nets")
        )
        for clinic in dataset["clinics"]
    )
    return {
        "metadata": dataset["metadata"],
        "assessments": assessments,
        "clinics": dataset["clinics"],
        "summary": {
            "zonesMonitored": len(assessments),
            "highRiskZones": sum(item["riskLevel"] in {"high", "critical"} for item in assessments),
            "facilitiesReporting": len(dataset["clinics"]),
            "supplyReadinessPct": int(_round(_average(readiness) * 100)),
            "currentCases": latest_cases,
            "previousCases": previous_cases,
            "caseChangePct": _round((latest_cases - previous_cases) / max(previous_cases, 1) * 100, 1),
            "tests": tests,
            "positiveTests": positives,
            "positivityPct": _round(positives / max(tests, 1) * 100, 1),
            "severeCases": sum(item["severeCases"] for item in clinical_context),
            "admissions": sum(item["admissions"] for item in clinical_context),
            "deaths": sum(item["deaths"] for item in clinical_context),
            "reportingCompletenessPct": _round(
                _average([item["reportingCompletenessPct"] for item in clinical_context]), 1
            ),
            "facilitiesBelowThreeWeeks": facilities_below_three_weeks,
            "stock": stock,
        },
    }


def run_response_cycle(
    dataset: dict[str, Any] | None = None, zone_id: str | None = None, now: str | None = None
) -> dict[str, Any]:
    dataset = dataset or load_dataset()
    created_at = now or utc_now()
    all_assessments = sorted(
        [assess_zone(zone, dataset["metadata"]) for zone in dataset["zones"]], key=lambda item: item["score"], reverse=True
    )
    requested = [item for item in all_assessments if item["zoneId"] == zone_id] if zone_id else all_assessments
    if not requested:
        raise ValueError(f"Unknown zone: {zone_id}")
    triggered = [item for item in requested if item["riskLevel"] in {"high", "critical"}]
    interventions = [
        {
            "zoneId": assessment["zoneId"],
            "zoneName": assessment["zoneName"],
            "assessment": assessment,
            "supplyPlan": plan_resources(dataset, assessment["zoneId"], assessment),
            "advisories": draft_advisories(dataset, assessment["zoneId"], assessment),
        }
        for assessment in triggered
    ]
    identifier_stamp = re.sub(r"[-:.TZ]", "", created_at)[:14]
    transfers = sum(len(item["supplyPlan"]["recommendations"]) for item in interventions)
    messages = sum(len(item["advisories"]["messages"]) for item in interventions)
    return {
        "id": f"cycle-{identifier_stamp}-{zone_id or 'network'}",
        "createdAt": created_at,
        "status": "pending_approval" if interventions else "monitoring",
        "mode": dataset["metadata"].get("mode", "unknown"),
        "assessments": requested,
        "interventions": interventions,
        "audit": [
            {"agent": "Surveillance", "status": "complete", "detail": f"{len(requested)} zone(s) assessed", "at": created_at},
            {"agent": "Supply", "status": "complete", "detail": f"{transfers} transfer drafts", "at": created_at},
            {"agent": "Community", "status": "complete", "detail": f"{messages} message drafts", "at": created_at},
            {
                "agent": "Human review",
                "status": "waiting" if interventions else "not_required",
                "detail": "Release and movement approval required" if interventions else "No action threshold crossed",
                "at": created_at,
            },
        ],
    }


def _read_runtime() -> dict[str, Any]:
    if not RUNTIME_PATH.exists():
        return {"cycles": []}
    try:
        runtime = _read_json(RUNTIME_PATH)
        return runtime if isinstance(runtime.get("cycles"), list) else {"cycles": []}
    except (json.JSONDecodeError, OSError):
        return {"cycles": []}


def create_cycle(zone_id: str | None = None) -> dict[str, Any]:
    runtime = _read_runtime()
    cycle = run_response_cycle(load_dataset(), zone_id=zone_id)
    runtime["cycles"] = [cycle, *runtime["cycles"]][:25]
    _atomic_write(RUNTIME_PATH, runtime)
    return cycle


def latest_cycle() -> dict[str, Any] | None:
    cycles = _read_runtime()["cycles"]
    return cycles[0] if cycles else None


def decide_cycle(cycle_id: str, decision: str, reviewer: str = "Operations reviewer") -> dict[str, Any]:
    if decision not in {"approve", "reject"}:
        raise ValueError("Decision must be approve or reject.")
    runtime = _read_runtime()
    cycle = next((item for item in runtime["cycles"] if item["id"] == cycle_id), None)
    if not cycle:
        raise ValueError("Response cycle was not found.")
    if cycle["status"] != "pending_approval":
        raise ValueError(f"Cycle is already {cycle['status']}.")
    cycle["status"] = "approved" if decision == "approve" else "rejected"
    cycle["review"] = {"decision": decision, "reviewer": reviewer, "at": utc_now()}
    cycle["audit"].append(
        {
            "agent": "Human review",
            "status": cycle["status"],
            "detail": f"{reviewer} {'approved' if decision == 'approve' else 'rejected'} the draft response",
            "at": cycle["review"]["at"],
        }
    )
    _atomic_write(RUNTIME_PATH, runtime)
    return cycle


def _finite_number(
    value: Any, field: str, *, minimum: float = 0, maximum: float = 9_007_199_254_740_991, integer: bool = False
) -> int | float:
    if isinstance(value, bool):
        raise ValueError(f"{field} must be a number.")
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{field} must be a number between {minimum} and {maximum}.") from error
    if not math.isfinite(number) or number < minimum or number > maximum or (integer and not number.is_integer()):
        kind = "a whole number" if integer else "a number"
        raise ValueError(f"{field} must be {kind} between {minimum} and {maximum}.")
    return int(number) if integer else number


def _stamp(dataset: dict[str, Any], source: str) -> dict[str, Any]:
    dataset["metadata"]["updatedAt"] = utc_now()
    dataset["metadata"]["mode"] = "operator-entry"
    dataset["metadata"]["lastUpdatedSource"] = source
    return dataset


def update_clinical(
    zone_id: str,
    weekly_cases: list[Any],
    tests: Any,
    positive_tests: Any,
    reporting_week: str,
    clinical_context: dict[str, Any] | None = None,
) -> dict[str, Any]:
    dataset = load_dataset()
    zone = next((item for item in dataset["zones"] if item["id"] == zone_id), None)
    if not zone:
        raise ValueError(f"Unknown zone: {zone_id}")
    if len(weekly_cases) != 6:
        raise ValueError("Weekly cases must contain exactly six totals.")
    normalized_cases = [_finite_number(value, f"Week {index + 1} cases", integer=True) for index, value in enumerate(weekly_cases)]
    normalized_tests = _finite_number(tests, "Tests", integer=True)
    normalized_positive = _finite_number(positive_tests, "Positive tests", integer=True)
    if normalized_positive > normalized_tests:
        raise ValueError("Positive tests cannot exceed total tests.")
    if not re.fullmatch(r"\d{4}-W(0[1-9]|[1-4]\d|5[0-3])", reporting_week):
        raise ValueError("Reporting week must use YYYY-Www, for example 2026-W41.")
    zone["weeklyCases"] = normalized_cases
    zone["tests"] = normalized_tests
    zone["positiveTests"] = normalized_positive
    if clinical_context is not None:
        normalized_context = {
            "suspectedCases": _finite_number(
                clinical_context.get("suspectedCases"), "Suspected cases", integer=True
            ),
            "severeCases": _finite_number(clinical_context.get("severeCases"), "Severe cases", integer=True),
            "admissions": _finite_number(clinical_context.get("admissions"), "Admissions", integer=True),
            "deaths": _finite_number(clinical_context.get("deaths"), "Deaths", integer=True),
            "under5ConfirmedCases": _finite_number(
                clinical_context.get("under5ConfirmedCases"), "Under-five confirmed cases", integer=True
            ),
            "pregnancyConfirmedCases": _finite_number(
                clinical_context.get("pregnancyConfirmedCases"), "Pregnancy confirmed cases", integer=True
            ),
            "chwFeverReferrals": _finite_number(
                clinical_context.get("chwFeverReferrals"), "CHW fever referrals", integer=True
            ),
            "reportingCompletenessPct": _finite_number(
                clinical_context.get("reportingCompletenessPct"),
                "Reporting completeness",
                maximum=100,
            ),
        }
        if normalized_tests > normalized_context["suspectedCases"]:
            raise ValueError("Tests performed cannot exceed suspected malaria cases.")
        if normalized_context["severeCases"] > normalized_positive:
            raise ValueError("Severe cases cannot exceed confirmed positive cases.")
        if normalized_context["admissions"] > normalized_context["severeCases"]:
            raise ValueError("Admissions cannot exceed severe cases.")
        if normalized_context["deaths"] > normalized_context["admissions"]:
            raise ValueError("Deaths cannot exceed malaria admissions.")
        if normalized_context["under5ConfirmedCases"] + normalized_context["pregnancyConfirmedCases"] > normalized_positive:
            raise ValueError("Priority-group confirmed cases cannot exceed all confirmed positive cases.")
        zone["clinical"] = normalized_context
    dataset["metadata"]["reportingWeek"] = reporting_week
    save_dataset(_stamp(dataset, "clinical-entry"))
    return zone


def update_environment(
    zone_id: str,
    rainfall: Any,
    temperature: Any,
    humidity: Any,
    rainfall_30d: Any | None = None,
    standing_water_index: Any | None = None,
    flooded_sites: Any | None = None,
) -> dict[str, Any]:
    dataset = load_dataset()
    zone = next((item for item in dataset["zones"] if item["id"] == zone_id), None)
    if not zone:
        raise ValueError(f"Unknown zone: {zone_id}")
    existing = zone.get("environment", {})
    zone["environment"] = {
        "rainfall7dMm": _finite_number(rainfall, "7-day rainfall", maximum=2_000),
        "rainfall30dMm": _finite_number(
            rainfall_30d if rainfall_30d is not None else existing.get("rainfall30dMm", 0),
            "30-day rainfall",
            maximum=4_000,
        ),
        "temperatureC": _finite_number(temperature, "Temperature", minimum=-10, maximum=60),
        "humidityPct": _finite_number(humidity, "Humidity", maximum=100),
        "standingWaterIndex": _finite_number(
            standing_water_index if standing_water_index is not None else existing.get("standingWaterIndex", 0),
            "Standing-water index",
            maximum=100,
        ),
        "floodedSites": _finite_number(
            flooded_sites if flooded_sites is not None else existing.get("floodedSites", 0),
            "Flooded sites",
            integer=True,
        ),
    }
    save_dataset(_stamp(dataset, "environment-entry"))
    return zone


def update_interventions(
    zone_id: str,
    coverage_date: str,
    llin_use_pct: Any,
    iptp3_coverage_pct: Any,
    vaccine_dose4_coverage_pct: Any,
    larval_sites_managed_pct: Any,
) -> dict[str, Any]:
    dataset = load_dataset()
    zone = next((item for item in dataset["zones"] if item["id"] == zone_id), None)
    if not zone:
        raise ValueError(f"Unknown zone: {zone_id}")
    if not _is_iso_date(coverage_date):
        raise ValueError("Coverage date must use YYYY-MM-DD.")
    zone["interventions"] = {
        "coverageDate": coverage_date,
        "llinUsePct": _finite_number(llin_use_pct, "LLIN use", maximum=100),
        "iptp3CoveragePct": _finite_number(iptp3_coverage_pct, "IPTp3 coverage", maximum=100),
        "vaccineDose4CoveragePct": _finite_number(
            vaccine_dose4_coverage_pct, "Malaria vaccine dose 4 coverage", maximum=100
        ),
        "larvalSitesManagedPct": _finite_number(
            larval_sites_managed_pct, "Larval sites managed", maximum=100
        ),
    }
    save_dataset(_stamp(dataset, "intervention-entry"))
    return zone


def update_supply(clinic_id: str, stock: dict[str, Any], weekly_use: dict[str, Any]) -> dict[str, Any]:
    dataset = load_dataset()
    clinic = next((item for item in dataset["clinics"] if item["id"] == clinic_id), None)
    if not clinic:
        raise ValueError(f"Unknown clinic: {clinic_id}")
    for group_name, values in (("stock", stock), ("weeklyUse", weekly_use)):
        clinic[group_name] = {
            resource: _finite_number(values.get(resource), f"{group_name} {resource.upper()}", integer=True)
            for resource in RESOURCE_LABELS
        }
    save_dataset(_stamp(dataset, "supply-entry"))
    return clinic


def _normalize_header(value: Any) -> str:
    return re.sub(r"[\s-]+", "_", str(value or "").lstrip("\ufeff").strip().lower())


def _number_field(
    row: dict[str, Any], key: str, errors: list[str], *, minimum: float = 0, maximum: float = 9_007_199_254_740_991, integer: bool = False
) -> int | float | None:
    value = row.get(key)
    if value in ("", None):
        errors.append(f"{key} is required.")
        return None
    try:
        return _finite_number(value, key, minimum=minimum, maximum=maximum, integer=integer)
    except ValueError as error:
        errors.append(str(error))
        return None


def _optional_number_field(
    row: dict[str, Any],
    key: str,
    errors: list[str],
    default: int | float,
    *,
    minimum: float = 0,
    maximum: float = 9_007_199_254_740_991,
    integer: bool = False,
) -> int | float | None:
    if key not in row or row.get(key) in ("", None):
        return default
    return _number_field(row, key, errors, minimum=minimum, maximum=maximum, integer=integer)


class _FirstHTMLTableParser(HTMLParser):
    """Extract cells from the first HTML table without executing or rendering markup."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.in_table = False
        self.finished = False
        self.in_cell = False
        self.current_row: list[str] | None = None
        self.current_cell: list[str] = []
        self.rows: list[list[str]] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        tag = tag.lower()
        if tag == "table" and not self.in_table and not self.finished:
            self.in_table = True
        elif self.in_table and tag == "tr":
            self.current_row = []
        elif self.in_table and tag in {"td", "th"} and self.current_row is not None:
            self.in_cell = True
            self.current_cell = []

    def handle_data(self, data: str) -> None:
        if self.in_cell:
            self.current_cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if self.in_table and tag in {"td", "th"} and self.in_cell and self.current_row is not None:
            self.current_row.append(" ".join("".join(self.current_cell).split()))
            self.current_cell = []
            self.in_cell = False
        elif self.in_table and tag == "tr" and self.current_row is not None:
            if any(cell.strip() for cell in self.current_row):
                self.rows.append(self.current_row)
            self.current_row = None
        elif self.in_table and tag == "table":
            self.in_table = False
            self.finished = True


def _records_from_rows(raw_rows: list[tuple[int, list[Any]]]) -> tuple[list[str], list[dict[str, Any]]]:
    if not raw_rows:
        return [], []
    columns = [_normalize_header(cell) for cell in raw_rows[0][1]]
    records = [
        {
            "row_number": row_number,
            "column_count": len(cells),
            "values": {
                column: str(cells[index]).strip() if index < len(cells) and cells[index] is not None else ""
                for index, column in enumerate(columns)
            },
        }
        for row_number, cells in raw_rows[1:]
    ]
    return columns, records


def _parse_uploaded(filename: str, content: bytes) -> tuple[str, list[str], list[dict[str, Any]]]:
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError("Upload files must be 15 MB or smaller.")
    extension = Path(filename).suffix.lower().lstrip(".")
    supported = {"csv", "tsv", "txt", "json", "xlsx", "html", "htm"}
    if extension not in supported:
        raise ValueError("Unsupported file format. Upload CSV, TSV, TXT, JSON, XLSX, or HTML.")

    if extension == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise ValueError("Excel support is unavailable. Install the dependencies in requirements.txt.") from error
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            worksheet = workbook.active
            raw_rows: list[tuple[int, list[Any]]] = []
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                cells = list(row)
                while cells and cells[-1] is None:
                    cells.pop()
                if any(cell is not None and str(cell).strip() for cell in cells):
                    raw_rows.append((row_number, cells))
            workbook.close()
        except Exception as error:
            raise ValueError("The Excel workbook could not be read. Upload a valid .xlsx file.") from error
        columns, records = _records_from_rows(raw_rows)
        return extension, columns, records

    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ValueError("The file must use UTF-8 text encoding.") from error
    if not text.strip():
        raise ValueError("Choose a non-empty file to upload.")

    if extension == "json":
        try:
            payload = json.loads(text)
        except json.JSONDecodeError as error:
            raise ValueError("The JSON file is not valid JSON.") from error
        values = payload if isinstance(payload, list) else payload.get("rows") if isinstance(payload, dict) else None
        if not isinstance(values, list):
            raise ValueError('JSON imports must be an array of rows or an object with a "rows" array.')
        if any(not isinstance(row, dict) for row in values):
            raise ValueError("Every JSON row must be an object.")
        records = [
            {
                "row_number": index + 1,
                "column_count": len(row),
                "values": {_normalize_header(key): value for key, value in row.items()},
            }
            for index, row in enumerate(values)
        ]
        columns = list(dict.fromkeys(key for record in records for key in record["values"]))
        return extension, columns, records

    if extension in {"html", "htm"}:
        parser = _FirstHTMLTableParser()
        try:
            parser.feed(text)
            parser.close()
        except Exception as error:
            raise ValueError("The HTML file could not be parsed.") from error
        if not parser.rows:
            raise ValueError("The HTML file does not contain a readable table.")
        columns, records = _records_from_rows(list(enumerate(parser.rows, start=1)))
        return "html", columns, records

    delimiter = "\t" if extension == "tsv" else ","
    if extension == "txt":
        try:
            delimiter = csv.Sniffer().sniff(text[:8_192], delimiters=",\t;|").delimiter
        except csv.Error:
            delimiter = ","
    try:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
        raw_rows = [(index, row) for index, row in enumerate(reader, start=1) if any(str(cell).strip() for cell in row)]
    except csv.Error as error:
        raise ValueError(f"The delimited file could not be parsed: {error}.") from error
    columns, records = _records_from_rows(raw_rows)
    return extension, columns, records


def _validate_row(kind: str, record: dict[str, Any], dataset: dict[str, Any]) -> tuple[dict[str, Any], list[str], str]:
    row = record["values"]
    errors: list[str] = []
    if kind == "clinical":
        zone_id = str(row.get("zone_id", "")).strip()
        zone = next((item for item in dataset["zones"] if item["id"] == zone_id), None)
        if not zone:
            errors.append(f"Unknown zone_id: {zone_id or '(blank)'}.")
        current = (zone or {}).get("clinical", {})
        reporting_week = str(row.get("reporting_week", "")).strip()
        if not re.fullmatch(r"\d{4}-W(0[1-9]|[1-4]\d|5[0-3])", reporting_week):
            errors.append("reporting_week must use YYYY-Www, for example 2026-W41.")
        weekly_cases = [_number_field(row, f"week_{index}_cases", errors, integer=True) for index in range(1, 7)]
        tests = _number_field(row, "tests", errors, integer=True)
        positive_tests = _number_field(row, "positive_tests", errors, integer=True)
        if tests is not None and positive_tests is not None and positive_tests > tests:
            errors.append("positive_tests cannot exceed tests.")
        clinical_context = {
            "suspectedCases": _optional_number_field(
                row,
                "suspected_cases",
                errors,
                max(current.get("suspectedCases", 0), tests or 0),
                integer=True,
            ),
            "severeCases": _optional_number_field(
                row, "severe_cases", errors, current.get("severeCases", 0), integer=True
            ),
            "admissions": _optional_number_field(
                row, "admissions", errors, current.get("admissions", 0), integer=True
            ),
            "deaths": _optional_number_field(row, "deaths", errors, current.get("deaths", 0), integer=True),
            "under5ConfirmedCases": _optional_number_field(
                row, "under_5_confirmed_cases", errors, current.get("under5ConfirmedCases", 0), integer=True
            ),
            "pregnancyConfirmedCases": _optional_number_field(
                row, "pregnancy_confirmed_cases", errors, current.get("pregnancyConfirmedCases", 0), integer=True
            ),
            "chwFeverReferrals": _optional_number_field(
                row, "chw_fever_referrals", errors, current.get("chwFeverReferrals", 0), integer=True
            ),
            "reportingCompletenessPct": _optional_number_field(
                row,
                "reporting_completeness_pct",
                errors,
                current.get("reportingCompletenessPct", 0),
                maximum=100,
            ),
        }
        suspected = clinical_context["suspectedCases"]
        severe = clinical_context["severeCases"]
        admissions = clinical_context["admissions"]
        deaths = clinical_context["deaths"]
        under_five = clinical_context["under5ConfirmedCases"]
        pregnancy = clinical_context["pregnancyConfirmedCases"]
        if "suspected_cases" in row and tests is not None and suspected is not None and tests > suspected:
            errors.append("tests cannot exceed suspected_cases.")
        if "severe_cases" in row and positive_tests is not None and severe is not None and severe > positive_tests:
            errors.append("severe_cases cannot exceed positive_tests.")
        if "admissions" in row and "severe_cases" in row and admissions is not None and severe is not None and admissions > severe:
            errors.append("admissions cannot exceed severe_cases.")
        if "deaths" in row and "admissions" in row and deaths is not None and admissions is not None and deaths > admissions:
            errors.append("deaths cannot exceed admissions.")
        if (
            "under_5_confirmed_cases" in row
            and "pregnancy_confirmed_cases" in row
            and positive_tests is not None
            and under_five is not None
            and pregnancy is not None
            and under_five + pregnancy > positive_tests
        ):
            errors.append("priority-group confirmed cases cannot exceed positive_tests.")
        normalized = {
            "zoneId": zone_id,
            "reportingWeek": reporting_week,
            "weeklyCases": weekly_cases,
            "tests": tests,
            "positiveTests": positive_tests,
            "clinical": clinical_context,
        }
        preview = (
            f"Week {reporting_week or '—'} · latest {weekly_cases[-1] if weekly_cases else '—'} · "
            f"{positive_tests}/{tests} positive · severe {severe}"
        )
    elif kind == "environment":
        zone_id = str(row.get("zone_id", "")).strip()
        zone = next((item for item in dataset["zones"] if item["id"] == zone_id), None)
        if not zone:
            errors.append(f"Unknown zone_id: {zone_id or '(blank)'}.")
        current = (zone or {}).get("environment", {})
        rainfall = _number_field(row, "rainfall_7d_mm", errors, maximum=2_000)
        temperature = _number_field(row, "temperature_c", errors, minimum=-10, maximum=60)
        humidity = _number_field(row, "humidity_pct", errors, maximum=100)
        rainfall_30d = _optional_number_field(
            row, "rainfall_30d_mm", errors, current.get("rainfall30dMm", 0), maximum=4_000
        )
        standing_water = _optional_number_field(
            row, "standing_water_index", errors, current.get("standingWaterIndex", 0), maximum=100
        )
        flooded_sites = _optional_number_field(
            row, "flooded_sites", errors, current.get("floodedSites", 0), integer=True
        )
        normalized = {
            "zoneId": zone_id,
            "rainfall7dMm": rainfall,
            "rainfall30dMm": rainfall_30d,
            "temperatureC": temperature,
            "humidityPct": humidity,
            "standingWaterIndex": standing_water,
            "floodedSites": flooded_sites,
        }
        preview = f"{rainfall} mm / 7d · {rainfall_30d} mm / 30d · standing water {standing_water}/100"
    elif kind == "intervention":
        zone_id = str(row.get("zone_id", "")).strip()
        if not any(zone["id"] == zone_id for zone in dataset["zones"]):
            errors.append(f"Unknown zone_id: {zone_id or '(blank)'}.")
        coverage_date = str(row.get("coverage_date", "")).strip()
        if not _is_iso_date(coverage_date):
            errors.append("coverage_date must use YYYY-MM-DD.")
        normalized = {
            "zoneId": zone_id,
            "coverageDate": coverage_date,
            "llinUsePct": _number_field(row, "llin_use_pct", errors, maximum=100),
            "iptp3CoveragePct": _number_field(row, "iptp3_coverage_pct", errors, maximum=100),
            "vaccineDose4CoveragePct": _number_field(
                row, "vaccine_dose4_coverage_pct", errors, maximum=100
            ),
            "larvalSitesManagedPct": _number_field(row, "larval_sites_managed_pct", errors, maximum=100),
        }
        preview = (
            f"LLIN {normalized['llinUsePct']}% · IPTp3 {normalized['iptp3CoveragePct']}% · "
            f"vaccine dose 4 {normalized['vaccineDose4CoveragePct']}%"
        )
    else:
        clinic_id = str(row.get("clinic_id", "")).strip()
        if not any(clinic["id"] == clinic_id for clinic in dataset["clinics"]):
            errors.append(f"Unknown clinic_id: {clinic_id or '(blank)'}.")
        stock = {resource: _number_field(row, f"stock_{resource}", errors, integer=True) for resource in RESOURCE_LABELS}
        weekly_use = {
            resource: _number_field(row, f"weekly_use_{resource}", errors, integer=True) for resource in RESOURCE_LABELS
        }
        normalized = {"clinicId": clinic_id, "stock": stock, "weeklyUse": weekly_use}
        preview = f"Stock RDT {stock['rdt']} · ACT {stock['act']} · nets {stock['nets']}"
    return normalized, errors, preview


def parse_data_import(kind: str, filename: str, content: bytes, dataset: dict[str, Any] | None = None) -> dict[str, Any]:
    if kind not in IMPORT_SCHEMAS:
        raise ValueError("Import type must be clinical, environment, intervention, or supply.")
    dataset = dataset or load_dataset()
    file_format, columns, records = _parse_uploaded(filename, content)
    schema = IMPORT_SCHEMAS[kind]
    file_errors: list[str] = []
    if not records:
        file_errors.append("The file does not contain any data rows.")
    duplicate_columns = sorted({column for column in columns if columns.count(column) > 1})
    if duplicate_columns:
        file_errors.append(f"Duplicate columns: {', '.join(duplicate_columns)}.")
    if "" in columns:
        file_errors.append("Every column must have a header name.")
    missing = [column for column in schema["required"] if column not in columns]
    if missing:
        file_errors.append(f"Missing required columns: {', '.join(missing)}.")
    sensitive = [column for column in columns if column in SENSITIVE_COLUMNS]
    if sensitive:
        file_errors.append(f"Patient-identifying columns are not accepted: {', '.join(sensitive)}. Upload aggregate data only.")

    rows: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    for record in records:
        normalized, errors, preview = _validate_row(kind, record, dataset)
        identifier = normalized.get("zoneId") or normalized.get("clinicId") or ""
        if record["column_count"] != len(columns):
            errors.insert(0, f"Expected {len(columns)} columns but found {record['column_count']}.")
        if identifier in seen_ids:
            errors.append(f"Duplicate {schema['id_field']}: {identifier}.")
        if identifier:
            seen_ids.add(identifier)
        rows.append(
            {
                "rowNumber": record["row_number"],
                "identifier": identifier,
                "valid": not errors,
                "errors": errors,
                "preview": preview,
                "normalized": normalized,
            }
        )
    if kind == "clinical":
        weeks = {row["normalized"]["reportingWeek"] for row in rows if row["valid"]}
        if len(weeks) > 1:
            file_errors.append("All clinical rows in one import must use the same reporting_week.")
    valid_count = sum(row["valid"] for row in rows)
    known_columns = [*schema["required"], *schema.get("optional", [])]
    warnings = [column for column in columns if column not in known_columns and column not in SENSITIVE_COLUMNS]
    return {
        "kind": kind,
        "filename": filename,
        "format": file_format,
        "columns": columns,
        "rowCount": len(rows),
        "validCount": valid_count,
        "invalidCount": len(rows) - valid_count,
        "canApply": bool(rows) and valid_count == len(rows) and not file_errors,
        "fileErrors": file_errors,
        "warnings": [f"Ignored columns: {', '.join(warnings)}."] if warnings else [],
        "rows": rows,
        "normalizedRows": [row["normalized"] for row in rows],
    }


def apply_parsed_import(dataset: dict[str, Any], result: dict[str, Any]) -> dict[str, Any]:
    if not result["canApply"]:
        raise ValueError("Resolve all import errors before applying the file.")
    updated = deepcopy(dataset)
    if result["kind"] == "clinical":
        for row in result["normalizedRows"]:
            zone = next(item for item in updated["zones"] if item["id"] == row["zoneId"])
            zone["weeklyCases"] = row["weeklyCases"]
            zone["tests"] = row["tests"]
            zone["positiveTests"] = row["positiveTests"]
            zone["clinical"] = row["clinical"]
            updated["metadata"]["reportingWeek"] = row["reportingWeek"]
    elif result["kind"] == "environment":
        for row in result["normalizedRows"]:
            zone = next(item for item in updated["zones"] if item["id"] == row["zoneId"])
            zone["environment"] = {
                "rainfall7dMm": row["rainfall7dMm"],
                "rainfall30dMm": row["rainfall30dMm"],
                "temperatureC": row["temperatureC"],
                "humidityPct": row["humidityPct"],
                "standingWaterIndex": row["standingWaterIndex"],
                "floodedSites": row["floodedSites"],
            }
    elif result["kind"] == "intervention":
        for row in result["normalizedRows"]:
            zone = next(item for item in updated["zones"] if item["id"] == row["zoneId"])
            zone["interventions"] = {
                "coverageDate": row["coverageDate"],
                "llinUsePct": row["llinUsePct"],
                "iptp3CoveragePct": row["iptp3CoveragePct"],
                "vaccineDose4CoveragePct": row["vaccineDose4CoveragePct"],
                "larvalSitesManagedPct": row["larvalSitesManagedPct"],
            }
    else:
        for row in result["normalizedRows"]:
            clinic = next(item for item in updated["clinics"] if item["id"] == row["clinicId"])
            clinic["stock"] = row["stock"]
            clinic["weeklyUse"] = row["weeklyUse"]
    return updated


def apply_data_import(result: dict[str, Any]) -> dict[str, Any]:
    updated = apply_parsed_import(load_dataset(), result)
    imported_at = utc_now()
    _stamp(updated, f"{result['kind']}-file-import")
    updated["metadata"]["lastImport"] = {
        "kind": result["kind"],
        "filename": result["filename"][:180],
        "rows": result["rowCount"],
        "importedAt": imported_at,
    }
    save_dataset(updated)
    return updated["metadata"]


def import_template(kind: str, dataset: dict[str, Any] | None = None) -> bytes:
    if kind not in IMPORT_SCHEMAS:
        raise ValueError("Unknown import type.")
    dataset = dataset or load_dataset()
    output = io.StringIO()
    fieldnames = [*IMPORT_SCHEMAS[kind]["required"], *IMPORT_SCHEMAS[kind].get("optional", [])]
    writer = csv.DictWriter(output, fieldnames=fieldnames, lineterminator="\n")
    writer.writeheader()
    if kind == "clinical":
        for zone in dataset["zones"]:
            row = {"zone_id": zone["id"], "reporting_week": dataset["metadata"]["reportingWeek"]}
            row.update({f"week_{index}_cases": value for index, value in enumerate(zone["weeklyCases"], start=1)})
            row.update({"tests": zone["tests"], "positive_tests": zone["positiveTests"]})
            clinical = zone["clinical"]
            row.update(
                {
                    "suspected_cases": clinical["suspectedCases"],
                    "severe_cases": clinical["severeCases"],
                    "admissions": clinical["admissions"],
                    "deaths": clinical["deaths"],
                    "under_5_confirmed_cases": clinical["under5ConfirmedCases"],
                    "pregnancy_confirmed_cases": clinical["pregnancyConfirmedCases"],
                    "chw_fever_referrals": clinical["chwFeverReferrals"],
                    "reporting_completeness_pct": clinical["reportingCompletenessPct"],
                }
            )
            writer.writerow(row)
    elif kind == "environment":
        for zone in dataset["zones"]:
            writer.writerow(
                {
                    "zone_id": zone["id"],
                    "rainfall_7d_mm": zone["environment"]["rainfall7dMm"],
                    "rainfall_30d_mm": zone["environment"]["rainfall30dMm"],
                    "temperature_c": zone["environment"]["temperatureC"],
                    "humidity_pct": zone["environment"]["humidityPct"],
                    "standing_water_index": zone["environment"]["standingWaterIndex"],
                    "flooded_sites": zone["environment"]["floodedSites"],
                }
            )
    elif kind == "intervention":
        for zone in dataset["zones"]:
            interventions = zone["interventions"]
            writer.writerow(
                {
                    "zone_id": zone["id"],
                    "coverage_date": interventions["coverageDate"],
                    "llin_use_pct": interventions["llinUsePct"],
                    "iptp3_coverage_pct": interventions["iptp3CoveragePct"],
                    "vaccine_dose4_coverage_pct": interventions["vaccineDose4CoveragePct"],
                    "larval_sites_managed_pct": interventions["larvalSitesManagedPct"],
                }
            )
    else:
        for clinic in dataset["clinics"]:
            writer.writerow(
                {
                    "clinic_id": clinic["id"],
                    "stock_rdt": clinic["stock"]["rdt"],
                    "stock_act": clinic["stock"]["act"],
                    "stock_nets": clinic["stock"]["nets"],
                    "weekly_use_rdt": clinic["weeklyUse"]["rdt"],
                    "weekly_use_act": clinic["weeklyUse"]["act"],
                    "weekly_use_nets": clinic["weeklyUse"]["nets"],
                }
            )
    return output.getvalue().encode("utf-8")
